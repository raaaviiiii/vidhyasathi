"""QA: INGESTION & STORAGE (src/ingest.py, src/chunk.py, src/build_index.py,
src/store.py, src/ocr.py failure behaviour).

Tests are written as they SHOULD pass per qa/SPEC.md. A failing test is a bug.
Real Qdrant server; fake embedder / OCR from tests/conftest.py.
"""
import hashlib
import math
import re
import threading
from pathlib import Path

import pytest

from src import config
from src.config import CHUNK_SIZE, CHUNK_OVERLAP

USER_A = "a@x.com"
USER_B = "b@x.com"


# ------------------------------------------------------------------ helpers
def _points(client, user=None):
    """All points (payload only) for a user, straight from Qdrant."""
    from src.store import user_filter
    out, offset = [], None
    if not client.collection_exists(config.COLLECTION_NAME):
        return out
    while True:
        pts, offset = client.scroll(collection_name=config.COLLECTION_NAME,
                                    scroll_filter=user_filter(user), with_payload=True,
                                    with_vectors=False, limit=256, offset=offset)
        out.extend(pts)
        if offset is None:
            return out


def _sources(client, user):
    return {p.payload["source"] for p in _points(client, user)}


def _write_txt(folder: Path, name: str, paras):
    folder.mkdir(parents=True, exist_ok=True)
    p = folder / name
    p.write_text("\n\n".join(paras), encoding="utf-8")
    return p


def _mixed_pdf(path: Path):
    """Page 1 digital text, page 2 a scanned-looking page (no text layer)."""
    import pymupdf
    doc = pymupdf.open()
    pg = doc.new_page()
    pg.insert_textbox(pymupdf.Rect(50, 50, 550, 800),
                      "The Euclidean algorithm computes the gcd by repeated division.",
                      fontsize=11)
    pg2 = doc.new_page()
    pg2.draw_rect(pymupdf.Rect(100, 100, 300, 200), color=(0, 0, 0), fill=(0.5, 0.5, 0.5))
    doc.save(str(path))
    return path


def _dense_embed(texts):
    """Dense, realistic-looking vectors (like MiniLM output) — JSON-heavy floats."""
    out = []
    for t in texts:
        h = hashlib.sha256(t.encode()).digest()
        v = [math.sin((h[i % 32] + 1) * (i + 1) * 0.6180339887) for i in range(config.EMBED_DIM)]
        n = math.sqrt(sum(x * x for x in v))
        out.append([x / n for x in v])
    return out


def _auth(app_client, user):
    """Sign `user` up (once per client) and return bearer headers; {} for no user.
    Identity now comes from a session token, never from a client-supplied field."""
    if not user:
        return {}
    cache = app_client.__dict__.setdefault("_tokens", {})
    if user not in cache:
        app_client.post("/api/signup", json={"name": user.split("@")[0], "email": user,
                                             "password": "pw123456"})
        r = app_client.post("/api/login", json={"email": user, "password": "pw123456"})
        cache[user] = r.json()["token"]
    return {"Authorization": f"Bearer {cache[user]}"}


def _upload(app_client, name, data: bytes, user=USER_A):
    return app_client.post("/api/upload", files={"file": (name, data)},
                           data={"user": user}, headers=_auth(app_client, user)).json()


def _material(app_client, user=USER_A):
    return app_client.get("/api/material", params={"user": user},
                          headers=_auth(app_client, user)).json()["docs"]


def _raise_ocr(png):
    raise RuntimeError("Could not resolve authentication method / network unreachable")


# ================================================================== CHUNKER
def test_chunk_short_text_single_chunk():
    from src.chunk import chunk_text
    assert chunk_text("hello world", CHUNK_SIZE, CHUNK_OVERLAP) == ["hello world"]


def test_chunk_exactly_chunk_size_is_one_chunk():
    from src.chunk import chunk_text
    t = ("abcd " * (CHUNK_SIZE // 5))[:CHUNK_SIZE]
    assert len(t) == CHUNK_SIZE
    assert chunk_text(t, CHUNK_SIZE, CHUNK_OVERLAP) == [t]


def test_chunk_no_spaces_terminates_and_bounded():
    from src.chunk import chunk_text
    t = "x" * 5000
    cs = chunk_text(t, CHUNK_SIZE, CHUNK_OVERLAP)
    assert 1 < len(cs) < 20
    assert all(0 < len(c) <= CHUNK_SIZE for c in cs)


def test_chunk_covers_every_word_and_respects_size():
    from src.chunk import chunk_text
    toks = [f"w{i:05d}" for i in range(1200)]
    cs = chunk_text(" ".join(toks), CHUNK_SIZE, CHUNK_OVERLAP)
    assert all(len(c) <= CHUNK_SIZE for c in cs)
    joined = "\n".join(cs)
    missing = [t for t in toks if not re.search(rf"\b{t}\b", joined)]
    assert missing == []


def test_chunk_consecutive_chunks_overlap():
    from src.chunk import chunk_text
    toks = [f"w{i:05d}" for i in range(1200)]
    cs = chunk_text(" ".join(toks), CHUNK_SIZE, CHUNK_OVERLAP)
    for prev, nxt in zip(cs, cs[1:]):
        # the head of the next chunk must come from the tail of the previous one
        head = nxt[:60]
        assert head in prev, (prev[-80:], head)


def test_chunk_no_redundant_tail_chunk():
    """A final chunk that is wholly contained in its predecessor is pure duplication
    (it wastes a retrieval slot and double-counts a passage)."""
    from src.chunk import chunk_text
    # Non-periodic text: a uniform string like "y"*1450 makes ANY two equal-length
    # windows identical, so substring containment can't tell redundancy apart.
    nospace = "".join(str(i) for i in range(600))
    for t in (nospace[:1450], nospace[:2000], " ".join(f"t{i:04d}" for i in range(290))):
        cs = chunk_text(t, CHUNK_SIZE, CHUNK_OVERLAP)
        for prev, nxt in zip(cs, cs[1:]):
            assert nxt not in prev, (len(t), [len(c) for c in cs])


def test_chunk_records_never_emits_empty_chunks():
    from src.chunk import chunk_records
    recs = [{"loc": "p1", "order": 1, "text": ""},
            {"loc": "p2", "order": 2, "text": "   "},
            {"loc": "p3", "order": 3, "text": "real text"}]
    out = chunk_records(recs, source="S")
    assert [c["text"] for c in out] == ["real text"]


def test_chunk_records_ids_unique_and_fields_kept():
    from src.chunk import chunk_records
    recs = [{"loc": f"p{i}", "order": i, "text": "word " * 400, "ocr": i == 2}
            for i in range(1, 4)]
    out = chunk_records(recs, source="Mod 1")
    ids = [c["id"] for c in out]
    assert len(ids) == len(set(ids))
    assert all(c["source"] == "Mod 1" for c in out)
    assert {c["loc"] for c in out} == {"p1", "p2", "p3"}
    assert all(c["ocr"] == (c["loc"] == "p2") for c in out)


def test_chunk_unicode_malayalam_text():
    from src.chunk import chunk_text
    t = " ".join(["ഗണിതശാസ്ത്രം"] * 300)
    cs = chunk_text(t, CHUNK_SIZE, CHUNK_OVERLAP)
    assert all(c and len(c) <= CHUNK_SIZE for c in cs)
    assert all("ഗണിതശാസ്ത്രം" in c for c in cs)


# ================================================================== LOADERS
def test_pdf_loader_pages(make_files):
    from src.ingest import load_document
    recs = load_document(make_files("pdf"))
    assert [r["loc"] for r in recs] == ["p1", "p2", "p3"]
    assert "Division Algorithm" in recs[0]["text"]
    assert not any(r.get("ocr") for r in recs)


def test_pdf_uppercase_extension_supported(make_files, fakes, tmp_path):
    from src.build_index import ingest_file
    p = make_files("pdf", name="LECTURE.PDF")
    r = ingest_file(p, user=USER_A)
    assert r["ok"] and r["chunks"] == 3


def test_scanned_pdf_goes_through_ocr_and_is_flagged(make_files, fakes):
    from src.build_index import ingest_file
    r = ingest_file(make_files("scanned_pdf"), user=USER_A)
    assert r["ok"] and r["ocr"] == 1
    assert len(fakes.ocr.calls) == 1
    pts = _points(fakes.client, USER_A)
    assert pts and all(p.payload["ocr"] is True for p in pts)


def test_image_goes_through_ocr_and_is_flagged(make_files, fakes):
    from src.build_index import ingest_file
    r = ingest_file(make_files("png"), user=USER_A)
    assert r["ok"] and r["ocr"] == 1
    assert _points(fakes.client, USER_A)[0].payload["loc"] == "image"


def test_docx_heading_locator_and_table(make_files):
    from src.ingest import load_document
    recs = load_document(make_files("docx"))
    body = [r for r in recs if r["loc"] != "table"]
    assert body and all(r["loc"] == "Number Theory" for r in body)
    assert any("gcd" in r["text"] and "greatest common divisor" in r["text"]
               for r in recs if r["loc"] == "table")


def test_pptx_slides_and_tables(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from src.ingest import load_document
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Primes"
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "Term", "Meaning"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "prime", "only divisors 1 and itself"
    p = tmp_path / "deck.pptx"
    prs.save(str(p))
    recs = load_document(p)
    assert recs[0]["loc"] == "slide 1"
    assert "only divisors 1 and itself" in recs[0]["text"]


def test_pptx_text_inside_grouped_shapes_is_captured(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from src.ingest import load_document
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    grp = s.shapes.add_group_shape()
    tb = grp.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    tb.text_frame.text = "Fermat little theorem: a^p is congruent to a mod p"
    p = tmp_path / "grouped.pptx"
    prs.save(str(p))
    recs = load_document(p)
    assert any("Fermat little theorem" in r["text"] for r in recs), recs


def test_xlsx_all_sheets_captured(tmp_path):
    from openpyxl import Workbook
    from src.ingest import load_document
    wb = Workbook()
    ws = wb.active
    ws.title = "Terms"
    ws.append(["Term", "Meaning"])
    ws.append(["gcd", "greatest common divisor"])
    ws2 = wb.create_sheet("Formulas")
    ws2.append(["Name", "Formula"])
    ws2.append(["division", "a = bq + r"])
    ws2.append(["zero", 0])
    p = tmp_path / "book.xlsx"
    wb.save(str(p))
    recs = load_document(p)
    locs = [r["loc"] for r in recs]
    assert "Terms row 1" in locs and "Formulas row 1" in locs
    assert any("a = bq + r" in r["text"] for r in recs)
    assert any("Formula: 0" in r["text"] for r in recs)   # numeric zero is not "empty"
    assert len({r["order"] for r in recs}) == len(recs)


def test_csv_basic(make_files):
    from src.ingest import load_document
    recs = load_document(make_files("csv"))
    assert recs[0] == {"loc": "row 1", "order": 1,
                       "text": "Term: gcd; Meaning: greatest common divisor"}


def test_csv_with_utf8_bom_header_is_clean(tmp_path):
    """Excel 'CSV UTF-8' always writes a BOM; it must not leak into the header."""
    from src.ingest import load_document
    p = tmp_path / "excel.csv"
    p.write_bytes("﻿Term,Meaning\ngcd,greatest common divisor\n".encode("utf-8"))
    recs = load_document(p)
    assert recs[0]["text"] == "Term: gcd; Meaning: greatest common divisor"


def test_csv_ragged_rows_keep_extra_cells(tmp_path):
    """A row with more cells than the header must not silently lose data."""
    from src.ingest import load_document
    p = tmp_path / "ragged.csv"
    p.write_text("Term,Meaning\nprime,two divisors,EXTRA-NOTE-KEPT\nshort\n")
    recs = load_document(p)
    text = " ".join(r["text"] for r in recs)
    assert "EXTRA-NOTE-KEPT" in text
    assert "short" in text


def test_txt_and_md_loaders(make_files):
    from src.ingest import load_document
    for k in ("txt", "md"):
        recs = load_document(make_files(k))
        assert len(recs) == 1 and "Euclidean algorithm" in recs[0]["text"]


def test_txt_utf16_with_bom_decodes(tmp_path):
    """Windows Notepad 'Unicode' saves UTF-16 LE with BOM."""
    from src.ingest import load_document
    p = tmp_path / "notepad.txt"
    p.write_text("The Euclidean algorithm computes the gcd.", encoding="utf-16")
    recs = load_document(p)
    assert recs and "Euclidean algorithm" in recs[0]["text"], recs


def test_source_label_normalises_name():
    from src.ingest import source_label
    assert source_label(Path("Mod_1  .1.pdf")) == "Mod 1 .1"
    assert source_label(Path("number_theory.docx")) == "number theory"


# ================================================================== ingest_file: errors
def test_ingest_empty_txt_clear_error(make_files, fakes):
    from src.build_index import ingest_file
    r = ingest_file(make_files("empty_txt"), user=USER_A)
    assert r["ok"] is False and r["chunks"] == 0 and "no extractable text" in r["error"]
    assert not fakes.client.collection_exists(config.COLLECTION_NAME) or \
        _points(fakes.client, USER_A) == []


def test_ingest_unsupported_clear_error(make_files, fakes):
    from src.build_index import ingest_file
    r = ingest_file(make_files("unsupported"), user=USER_A)
    assert r["ok"] is False and "unsupported" in r["error"]


def test_ingest_corrupt_pdf_returns_error_not_exception(make_files, fakes):
    from src.build_index import ingest_file
    r = ingest_file(make_files("corrupt_pdf"), user=USER_A)   # must not raise
    assert r["ok"] is False and r.get("error")


def test_ingest_corrupt_docx_returns_error_not_exception(tmp_path, fakes):
    from src.build_index import ingest_file
    p = tmp_path / "broken.docx"
    p.write_bytes(b"PK\x03\x04 not really a zip")
    r = ingest_file(p, user=USER_A)
    assert r["ok"] is False and r.get("error")


def test_upload_corrupt_file_leaves_no_orphan(app_client, make_files):
    data = make_files("corrupt_pdf").read_bytes()
    r = _upload(app_client, "broken.pdf", data)
    assert r["indexed"] is False
    assert [d["name"] for d in _material(app_client)] == []


def test_upload_unsupported_and_empty_leave_no_orphan(app_client, make_files):
    r1 = _upload(app_client, "virus.exe", make_files("unsupported").read_bytes())
    r2 = _upload(app_client, "blank.txt", make_files("empty_txt").read_bytes())
    assert r1["indexed"] is False and r2["indexed"] is False
    assert _material(app_client) == []


def test_upload_filename_cannot_escape_user_folder(app_client):
    """The uploaded file must land inside the user's own folder, whatever the
    client claims the filename is."""
    _upload(app_client, "../../escape_probe.txt", b"The gcd is the greatest common divisor.")
    hits = list(app_client.tmp.rglob("escape_probe.txt"))
    up_root = app_client.tmp / "data" / "uploads"
    assert all(up_root in h.parents and h.parent.parent == up_root for h in hits), hits


# ================================================================== OCR failure
def test_pdf_with_scanned_page_degrades_when_ocr_fails(fakes, monkeypatch, tmp_path):
    """OCR down (no API key / no network) must not sink the whole upload: the
    digital pages still get indexed."""
    import src.ocr
    monkeypatch.setattr(src.ocr, "ocr_image", _raise_ocr)
    from src.build_index import ingest_file
    p = _mixed_pdf(tmp_path / "mixed.pdf")
    r = ingest_file(p, user=USER_A)               # must not raise
    assert r["ok"] is True and r["chunks"] >= 1
    assert any("Euclidean" in pt.payload["text"] for pt in _points(fakes.client, USER_A))


def test_image_when_ocr_fails_returns_error_not_exception(fakes, monkeypatch, make_files):
    import src.ocr
    monkeypatch.setattr(src.ocr, "ocr_image", _raise_ocr)
    from src.build_index import ingest_file
    r = ingest_file(make_files("png"), user=USER_A)   # must not raise
    assert r["ok"] is False and r.get("error")


def test_upload_image_when_ocr_fails_leaves_no_orphan(app_client, monkeypatch, make_files):
    import src.ocr
    monkeypatch.setattr(src.ocr, "ocr_image", _raise_ocr)
    r = _upload(app_client, "board.png", make_files("png").read_bytes())
    assert r["indexed"] is False
    assert _material(app_client) == []


# ================================================================== store: replace / isolation
def test_ingest_count_matches_chunks(make_files, fakes):
    from src.build_index import ingest_file
    from src.store import count
    r = ingest_file(make_files("pdf"), user=USER_A)
    assert {k: r[k] for k in ("ok", "source", "chunks", "ocr")} == \
        {"ok": True, "source": "sample pdf", "chunks": 3, "ocr": 0}
    assert count(USER_A) == 3


def test_payload_fields_complete(make_files, fakes):
    from src.build_index import ingest_file
    ingest_file(make_files("pdf"), user=USER_A)
    for p in _points(fakes.client, USER_A):
        pl = p.payload
        assert set(pl) >= {"chunk_id", "user", "source", "loc", "order", "text", "ocr"}
        assert pl["user"] == USER_A and pl["source"] == "sample pdf"


def test_reupload_same_file_replaces_not_duplicates(make_files, fakes):
    from src.build_index import ingest_file
    from src.store import count
    p = make_files("pdf")
    ingest_file(p, user=USER_A)
    ids1 = sorted(x.payload["chunk_id"] for x in _points(fakes.client, USER_A))
    ingest_file(p, user=USER_A)
    ids2 = sorted(x.payload["chunk_id"] for x in _points(fakes.client, USER_A))
    # Point ids are versioned (new copy written before the old one is removed), so
    # compare chunk identities: same chunks, no duplicates.
    assert count(USER_A) == 3 and ids1 == ids2


def test_reupload_shorter_version_drops_stale_chunks(fakes, tmp_path):
    from src.build_index import ingest_file
    from src.store import count
    d = tmp_path / "v"
    big = _write_txt(d / "1", "notes.txt", ["alpha beta gamma " * 300])
    r1 = ingest_file(big, user=USER_A)
    assert r1["chunks"] > 3
    small = _write_txt(d / "2", "notes.txt", ["just one short line"])
    ingest_file(small, user=USER_A)
    assert count(USER_A) == 1
    assert _points(fakes.client, USER_A)[0].payload["text"] == "just one short line"


def test_failed_reupload_keeps_previous_version(fakes, tmp_path, monkeypatch):
    """Replace must be safe: if storing the new version fails, the student must not
    be left with NEITHER version."""
    import src.build_index as bi
    from src.store import count
    p = _write_txt(tmp_path, "notes.txt", ["The gcd is the greatest common divisor."])
    bi.ingest_file(p, user=USER_A)
    assert count(USER_A) == 1

    def boom(*a, **k):
        raise RuntimeError("qdrant write failed")
    monkeypatch.setattr(bi, "upsert_chunks", boom)
    p.write_text("A revised version of the notes about primes.")
    with pytest.raises(RuntimeError):
        bi.ingest_file(p, user=USER_A)
    assert count(USER_A) == 1, "old chunks were deleted before the new ones were stored"


def test_same_file_two_users_isolated(make_files, fakes):
    from src.build_index import ingest_file
    from src.store import count, delete_user
    p = make_files("pdf")
    ingest_file(p, user=USER_A)
    ingest_file(p, user=USER_B)
    ida = {str(x.id) for x in _points(fakes.client, USER_A)}
    idb = {str(x.id) for x in _points(fakes.client, USER_B)}
    assert count(USER_A) == 3 and count(USER_B) == 3 and not (ida & idb)
    delete_user(USER_A)
    assert count(USER_A) == 0 and count(USER_B) == 3


def test_delete_user_blank_is_noop(make_files, fakes):
    from src.build_index import ingest_file
    from src.store import count, delete_user
    ingest_file(make_files("pdf"), user=USER_A)
    delete_user("")
    delete_user(None)
    assert count(USER_A) == 3


def test_delete_source_only_that_source(fakes, tmp_path):
    from src.build_index import ingest_file
    from src.store import delete_source
    ingest_file(_write_txt(tmp_path, "algebra.txt", ["groups rings fields"]), user=USER_A)
    ingest_file(_write_txt(tmp_path, "calculus.txt", ["limits derivatives"]), user=USER_A)
    ingest_file(_write_txt(tmp_path, "algebra.txt", ["groups rings fields"]), user=USER_B)
    delete_source(USER_A, "algebra")
    assert _sources(fakes.client, USER_A) == {"calculus"}
    assert _sources(fakes.client, USER_B) == {"algebra"}


def test_same_stem_different_extension_both_kept(make_files, fakes):
    """notes.pdf and notes.docx are two different files: SPEC 'Two different files
    must never overwrite each other's chunks.'"""
    from src.build_index import ingest_file
    from src.store import count
    ingest_file(make_files("pdf", name="notes.pdf"), user=USER_A)
    n_pdf = count(USER_A)
    r = ingest_file(make_files("docx", name="notes.docx"), user=USER_A)
    assert count(USER_A) == n_pdf + r["chunks"]


def test_remove_one_same_stem_file_keeps_the_other(app_client, make_files):
    _upload(app_client, "notes.pdf", make_files("pdf", name="notes.pdf").read_bytes())
    _upload(app_client, "notes.txt", make_files("txt", name="notes.txt").read_bytes())
    app_client.post("/api/material/delete", json={"user": USER_A, "name": "notes.pdf"},
                    headers=_auth(app_client, USER_A))
    docs = _material(app_client)
    assert [d["name"] for d in docs] == ["notes.txt"]
    assert docs[0]["indexed"] is True
    from src.store import count
    assert count(USER_A) >= 1


@pytest.mark.parametrize("n1,n2", [("Notes 1.txt", "notes-1.txt"),
                                   ("Notes.txt", "notes.txt"),
                                   ("C++ basics.txt", "C basics.txt")])
def test_names_that_slugify_alike_do_not_collide(fakes, tmp_path, n1, n2):
    from src.build_index import ingest_file
    from src.store import count
    a = _write_txt(tmp_path / "a", n1, ["First document about primes."])
    b = _write_txt(tmp_path / "b", n2, ["Second document about the gcd."])
    ingest_file(a, user=USER_A)
    ingest_file(b, user=USER_A)
    assert count(USER_A) == 2
    texts = {p.payload["text"] for p in _points(fakes.client, USER_A)}
    assert texts == {"First document about primes.", "Second document about the gcd."}


def test_non_ascii_filenames_do_not_collide(fakes, tmp_path):
    """Malayalam / Hindi file names all slugify to 'doc' -> same point ids."""
    from src.build_index import ingest_file
    from src.store import count
    ingest_file(_write_txt(tmp_path, "ഗണിതം.txt", ["Malayalam maths notes."]), user=USER_A)
    ingest_file(_write_txt(tmp_path, "भौतिकी.txt", ["Hindi physics notes."]), user=USER_A)
    assert count(USER_A) == 2
    assert _sources(fakes.client, USER_A) == {"ഗണിതം", "भौतिकी"}


# ================================================================== store: missing collection
def test_count_when_collection_missing_is_zero(fakes):
    from src.store import count
    assert not fakes.client.collection_exists(config.COLLECTION_NAME)
    assert count(USER_A) == 0
    assert count() == 0


def test_user_sources_and_delete_when_collection_missing(fakes):
    from src.store import user_sources, delete_source, delete_user
    assert user_sources(USER_A) == set()
    delete_source(USER_A, "x")
    delete_user(USER_A)


def test_retrieve_when_collection_missing_is_empty(fakes):
    from src.retrieve import retrieve
    assert retrieve("what is gcd", user=USER_A) == []


def test_material_list_before_any_upload(app_client):
    assert _material(app_client) == []


# ================================================================== store: scale
def test_user_sources_paginates_past_256(fakes):
    from src.store import upsert_chunks, user_sources
    chunks = [{"id": f"s{i}_1_c0", "source": f"src {i}", "text": f"t{i}"} for i in range(300)]
    upsert_chunks(chunks, fakes.embedder.embed_texts([c["text"] for c in chunks]), user=USER_A)
    assert len(user_sources(USER_A)) == 300


def test_upsert_empty_list_is_noop(fakes):
    from src.store import upsert_chunks, count
    upsert_chunks([], [], user=USER_A)
    assert count(USER_A) == 0


def test_large_textbook_ingests(fakes, tmp_path, monkeypatch):
    """A ~600-page textbook (~5000 chunks) with real-looking dense vectors must
    index (upserts must be batched under Qdrant's 32 MB request limit)."""
    import src.build_index as bi
    monkeypatch.setattr(bi, "embed_texts", _dense_embed)
    from src.store import count
    para = " ".join(f"term{i % 997} definition example" for i in range(130))
    p = _write_txt(tmp_path, "textbook.txt", [para] * 950)
    r = bi.ingest_file(p, user=USER_A)
    assert r["ok"] and r["chunks"] > 4000
    assert count(USER_A) == r["chunks"]


# ================================================================== concurrency
def _ingest_many(tmp_path, n, user_of):
    from src.build_index import ingest_file
    paths = [_write_txt(tmp_path / f"d{i}", f"doc{i}.txt",
                        [f"Document {i} paragraph {j} " * 40 for j in range(4)])
             for i in range(n)]
    results, errors = [None] * n, []

    def work(i):
        try:
            results[i] = ingest_file(paths[i], user=user_of(i))
        except Exception as e:      # noqa: BLE001
            errors.append(f"{i}: {type(e).__name__}: {str(e)[:120]}")
    ts = [threading.Thread(target=work, args=(i,)) for i in range(n)]
    for t in ts:
        t.start()
    for t in ts:
        t.join()
    return results, errors


def test_concurrent_first_uploads_on_fresh_collection(fakes, tmp_path):
    """Several first-ever uploads at once (fresh install, threadpool) must all
    succeed — no 'collection already exists' race."""
    from src.store import count
    assert not fakes.client.collection_exists(config.COLLECTION_NAME)
    results, errors = _ingest_many(tmp_path, 8, lambda i: f"u{i}@x.com")
    assert errors == []
    for i, r in enumerate(results):
        assert count(f"u{i}@x.com") == r["chunks"]


def test_concurrent_uploads_no_lost_writes(fakes, tmp_path):
    from src.store import reset_collection, count, user_sources
    reset_collection()
    counted = []
    stop = threading.Event()

    def reader():
        while not stop.is_set():
            counted.append(count(USER_A))
    rt = threading.Thread(target=reader)
    rt.start()
    results, errors = _ingest_many(tmp_path, 12, lambda i: USER_A if i % 2 else USER_B)
    stop.set()
    rt.join()
    assert errors == []
    exp_a = sum(r["chunks"] for i, r in enumerate(results) if i % 2)
    exp_b = sum(r["chunks"] for i, r in enumerate(results) if not i % 2)
    assert count(USER_A) == exp_a and count(USER_B) == exp_b
    assert user_sources(USER_A) == {f"doc{i}" for i in range(12) if i % 2}
    assert counted and max(counted) <= exp_a


def test_concurrent_double_upload_same_file_no_duplicates(fakes, tmp_path):
    from src.build_index import ingest_file
    from src.store import reset_collection, count
    reset_collection()
    p = _write_txt(tmp_path, "notes.txt", ["The gcd paragraph. " * 100])
    errs = []

    def work():
        try:
            ingest_file(p, user=USER_A)
        except Exception as e:      # noqa: BLE001
            errs.append(e)
    ts = [threading.Thread(target=work) for _ in range(6)]
    for t in ts:
        t.start()
    for t in ts:
        t.join()
    assert errs == []
    single = ingest_file(p, user=USER_A)["chunks"]
    assert count(USER_A) == single


# ================================================================== build()
def _patch_build(monkeypatch, docs):
    import src.build_index as bi
    monkeypatch.setattr(bi, "list_documents", lambda: docs)
    monkeypatch.setattr(bi, "unsupported_files", lambda: [])
    monkeypatch.setattr(bi, "close", lambda: None)   # keep the shared test client open
    return bi


def test_build_indexes_all_documents(fakes, make_files, monkeypatch):
    docs = [make_files("pdf"), make_files("docx"), make_files("csv")]
    bi = _patch_build(monkeypatch, docs)
    from src.store import count, user_sources
    bi.build(owner="seed@x.com")
    assert user_sources("seed@x.com") == {"sample pdf", "sample docx", "sample csv"}
    assert count("seed@x.com") >= 3 + 2 + 2


def test_build_does_not_wipe_students_uploads(fakes, make_files, monkeypatch, tmp_path):
    """SPEC: uploaded material persists. Re-seeding the shared KB from data/docs must
    not erase every student's uploads."""
    from src.build_index import ingest_file
    from src.store import count
    ingest_file(_write_txt(tmp_path, "mine.txt", ["My own notes on primes."]), user=USER_A)
    bi = _patch_build(monkeypatch, [make_files("pdf")])
    bi.build(owner=None)
    assert count(USER_A) == 1


# ================================================================== cross-user / traversal data loss
def test_delete_source_without_user_does_not_touch_other_students(fakes, tmp_path):
    """delete_source(None, ...) (blank / shared user) must never wipe a named
    student's chunks for a same-named document."""
    from src.build_index import ingest_file
    from src.store import count
    ingest_file(_write_txt(tmp_path / "a", "notes.txt", ["Student A notes on primes."]), user=USER_A)
    ingest_file(_write_txt(tmp_path / "b", "notes.txt", ["Student B notes on gcd."]), user=USER_B)
    ingest_file(_write_txt(tmp_path / "s", "notes.txt", ["Shared notes."]), user=None)
    assert count(USER_A) == 1 and count(USER_B) == 1


def test_upload_with_blank_user_does_not_delete_other_students_material(app_client, make_files):
    _upload(app_client, "notes.txt", b"Student A notes on the Euclidean algorithm.", user=USER_A)
    _upload(app_client, "notes.txt", b"Anonymous notes.", user="")
    docs = _material(app_client, USER_A)
    assert [d["name"] for d in docs] == ["notes.txt"]
    assert docs[0]["indexed"] is True, "student A's chunks were deleted by an anonymous upload"


def test_upload_traversal_name_cannot_overwrite_or_delete_app_files(app_client):
    """A crafted filename must not be able to clobber (and then unlink, via the
    'unsupported -> delete orphan' path) files outside the user's folder,
    e.g. the accounts database."""
    app_client.post("/api/signup", json={"name": "A", "email": USER_A, "password": "pw1234"})
    _auth(app_client, USER_B)          # sign the attacker in first (a legitimate DB write)
    db = app_client.tmp / "users.db"
    assert db.exists()
    before = db.read_bytes()
    _upload(app_client, "../../../users.db", b"garbage", user=USER_B)
    assert db.exists() and db.read_bytes() == before
    r = app_client.post("/api/login", json={"email": USER_A, "password": "pw1234"})
    assert r.status_code == 200
