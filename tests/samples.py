"""Real sample documents for every supported format, built on the fly."""
from pathlib import Path

NUMBER_THEORY = [
    "The Division Algorithm states that for integers a and b with b > 0 there exist "
    "unique integers q and r such that a = bq + r and 0 <= r < b. q is the quotient and "
    "r is the remainder.",
    "The greatest common divisor gcd(a, b) is the largest integer dividing both a and b. "
    "The Euclidean algorithm computes the gcd by repeated division.",
    "A prime number is an integer greater than 1 whose only positive divisors are 1 and "
    "itself. The Fundamental Theorem of Arithmetic says every integer factors uniquely "
    "into primes.",
]


def make_sample(kind: str, folder: Path, name: str | None = None, paras=None) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    paras = paras or NUMBER_THEORY
    ext = {"pdf": "pdf", "scanned_pdf": "pdf", "docx": "docx", "pptx": "pptx",
           "xlsx": "xlsx", "csv": "csv", "txt": "txt", "md": "md", "png": "png",
           "empty_txt": "txt", "corrupt_pdf": "pdf", "unsupported": "exe"}[kind]
    path = folder / (name or f"sample_{kind}.{ext}")

    if kind == "pdf":
        import pymupdf
        doc = pymupdf.open()
        for p in paras:
            page = doc.new_page()
            page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), p, fontsize=11)
        doc.save(str(path))
    elif kind == "scanned_pdf":
        import pymupdf
        doc = pymupdf.open()
        page = doc.new_page()
        page.draw_rect(pymupdf.Rect(100, 100, 300, 200), color=(0, 0, 0), fill=(0.5, 0.5, 0.5))
        doc.save(str(path))
    elif kind == "docx":
        from docx import Document
        d = Document()
        d.add_heading("Number Theory", level=1)
        for p in paras:
            d.add_paragraph(p)
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Term", "Meaning"
        t.cell(1, 0).text, t.cell(1, 1).text = "gcd", "greatest common divisor"
        d.save(str(path))
    elif kind == "pptx":
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        for p in paras:
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = "Slide"
            tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
            tb.text_frame.text = p
        prs.save(str(path))
    elif kind == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Terms"
        ws.append(["Term", "Meaning"])
        ws.append(["gcd", "greatest common divisor of two integers"])
        ws.append(["prime", "integer greater than 1 with only divisors 1 and itself"])
        wb.save(str(path))
    elif kind == "csv":
        path.write_text("Term,Meaning\ngcd,greatest common divisor\n"
                        "prime,integer with exactly two positive divisors\n")
    elif kind in ("txt", "md"):
        path.write_text("\n\n".join(paras))
    elif kind == "png":
        from PIL import Image, ImageDraw
        img = Image.new("RGB", (400, 120), "white")
        ImageDraw.Draw(img).text((10, 40), "a = bq + r", fill="black")
        img.save(str(path))
    elif kind == "empty_txt":
        path.write_text("   \n  ")
    elif kind == "corrupt_pdf":
        path.write_bytes(b"%PDF-1.4 this is not really a pdf \x00\x01\x02")
    elif kind == "unsupported":
        path.write_bytes(b"MZ\x90\x00binary")
    return path
